# Copyright (c) 2026, Sidharth PV and contributors
# For license information, please see license.txt

"""Export any DocType as a three-sheet workbook.

Sheet 1 (named after the DocType) lists every field with its database type, whether it is
required, a description and an example. Sheet 2 ("Payload") is an example record as JSON and
sheet 3 ("Schema") is a JSON Schema for that payload. Child tables are flattened into sheet 1
as ``fieldname (child: DocType)`` rows and nested as arrays in the payload and schema.

Everything is read live from the DocType meta and the database, so it is always current.
Nothing is stored: there are no tables or records behind this.
"""

import io
import json
import re
import zipfile

import frappe
from frappe import _
from frappe.model import table_fields
from frappe.utils import add_days, cint, strip_html_tags, today
from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Font, PatternFill

# Layout-only fieldtypes never hold data. Image just points at another Attach field.
SKIP_FIELDTYPES = {"Section Break", "Column Break", "Tab Break", "HTML", "Heading", "Button", "Fold", "Image"}
HEADERS = ("Field Name", "Data Type", "Required", "Description", "Example")
COLUMN_WIDTHS = (34, 16, 10, 95, 34)
TEXT_SHEET_WIDTH = 120
CHILD_TABLE_TYPE = "N/A (child table)"
INVALID_SHEET_CHARS = re.compile(r"[\[\]:*?/\\]")
MAX_CELL_CHARS = 32000  # Excel allows 32,767 per cell

# Defaults Frappe resolves at runtime, in words a reader understands
DYNAMIC_DEFAULTS = {
	"Today": "today's date",
	"Now": "the current time",
	"__user": "the current user",
	"user": "the current user",
}

# Example values that keep a sample record believable. Names are matched as whole words.
LATER_DATE = re.compile(
	r"(^|_)(due|delivery|expected|valid|expiry|expire|end|to|required|schedule|release|lead)(_|$)"
)
ZERO_AMOUNT = re.compile(
	r"(^|_)(paid|advance|advances|write_off|writeoff|discount|rounding|adjustment|withheld|allocated|"
	r"unallocated|tax|taxes|charges|commission|incentives|rejected)(_|$)"
)
MAX_PACK = 50  # DocTypes per multi-DocType download, to bound the time of one request
STANDARD_FIELDS_PARENT = ("owner", "creation", "modified", "modified_by", "docstatus", "idx")
STANDARD_FIELDS_CHILD = ("parent", "parentfield", "parenttype", "idx")
RATE_OF_ONE = re.compile(r"(^|_)(conversion|exchange)_rate$")
QUANTITY = re.compile(r"(^|_)(qty|quantity)(_|$)")
PROGRESS = re.compile(r"^per_")


# ---- field collection -------------------------------------------------------------------


def get_column_types(doctype):
	"""Real database column types for a DocType's table, e.g. {"grand_total": "decimal(21,9)"}."""
	if frappe.db.db_type != "mariadb":
		# Other engines (Postgres, SQLite): let Frappe describe the table. Not tested here, where
		# only MariaDB is available. The MariaDB query below is kept because Frappe's own helper does
		# not restrict itself to the current database, which matters on a server hosting several sites.
		try:
			return {
				col["name"]: col["type"] for col in frappe.db.get_table_columns_description("tab" + doctype)
			}
		except Exception:
			return {}
	rows = frappe.db.sql(
		"""select column_name, column_type from information_schema.columns
		where table_schema = database() and table_name = %s""",
		("tab" + doctype,),
	)
	return {name: column_type for name, column_type in rows}


def type_from_fieldtype(fieldtype):
	"""Database type Frappe would create for a fieldtype, e.g. Currency -> decimal(21,9).

	Used where there is no real table to read: Single and virtual DocTypes.
	"""
	entry = frappe.db.type_map.get(fieldtype)
	if not entry or not entry[0]:
		return None
	name, length = entry
	return f"{name}({length})" if length else name


NUMBER_FIELDTYPES = {"Currency", "Float", "Percent", "Rating", "Duration"}
INTEGER_FIELDTYPES = {"Int", "Long Int", "Check"}


def json_type(db_type, fieldtype=None):
	"""JSON Schema type of a column.

	The database type decides when it is informative. Some engines (SQLite) report every column as
	TEXT, so then the Frappe fieldtype decides: a Currency is a number whatever the engine calls it.
	"""
	db_type = (db_type or "").lower()
	if db_type.startswith(("decimal", "float", "double", "numeric", "real")):
		return "number"
	if db_type.startswith(("int", "tinyint", "smallint", "mediumint", "bigint")):
		return "integer"
	if fieldtype in NUMBER_FIELDTYPES:
		return "number"
	if fieldtype in INTEGER_FIELDTYPES:
		return "integer"
	return "string"


def collect_fields(doctype, is_child=False, include_standard=False, context=None, customisations_only=False):
	"""Ordered field records for one DocType: required first, then optional, both in meta order.

	Hidden, virtual, layout and column-less fields are left out. Table fields are kept. With
	`include_standard`, the system-maintained columns (owner, docstatus, parent...) come last.
	`context` says which parent and table a child row belongs to, for their examples. With
	`customisations_only`, only fields that a Custom Field or Property Setter added or changed are
	kept (every field of a custom DocType counts), and the generated `name` row is left out.
	"""
	meta = frappe.get_meta(doctype)
	column_types = get_column_types(doctype)
	# Single and virtual DocTypes have no table of their own to read column types from
	has_table = bool(column_types) and not cint(meta.issingle) and not cint(meta.get("is_virtual"))
	custom = custom_sources(doctype)
	masters = default_masters()
	whole = bool(cint(meta.custom))  # a custom DocType is customisation from top to bottom
	customised = (
		set(custom) | set(property_setter_fields(doctype)) if customisations_only and not whole else set()
	)

	records = []
	if not is_child and has_table and "name" in column_types and (whole or not customisations_only):
		records.append(
			{
				"fieldname": "name",
				"df": frappe._dict(fieldname="name", label="Name", fieldtype="Data", options=None),
				"db_type": column_types["name"],
				"required": True,
				"is_table": False,
				"description": name_description(meta),
				"example": name_example(meta, column_types["name"]),
			}
		)

	optional, required = [], []
	for df in meta.fields:
		if df.fieldtype in SKIP_FIELDTYPES or cint(df.hidden) or cint(df.get("is_virtual")):
			continue
		is_table = df.fieldtype in table_fields
		if customisations_only and not whole and not is_table and df.fieldname not in customised:
			continue
		if is_table:
			db_type = CHILD_TABLE_TYPE
		elif has_table:
			if df.fieldname not in column_types:
				continue
			db_type = column_types[df.fieldname]
		else:
			db_type = type_from_fieldtype(df.fieldtype)
			if not db_type:
				continue

		record = {
			"fieldname": df.fieldname,
			"df": df,
			"db_type": db_type,
			"required": bool(cint(df.reqd)),
			"conditional": bool(df.mandatory_depends_on) and not cint(df.reqd),
			"is_table": is_table,
			"child_doctype": df.options if is_table else None,
			"customised": whole or df.fieldname in customised or not customisations_only,
			"description": describe(df, custom.get(df.fieldname), doctype),
		}
		if not is_table:
			record["example"] = example_value(df, record["db_type"], masters)
		(required if record["required"] else optional).append(record)

	standard = (
		standard_fields(meta, column_types, is_child, context) if include_standard and has_table else []
	)
	return records + required + optional + standard


def standard_fields(meta, column_types, is_child, context):
	"""The columns every table has and the system maintains. Integrators read them, never send them."""
	context = context or {}
	moment = f"{today()} 10:00:00.000000"
	info = {
		"owner": ("User who created the record.", "Administrator", "Data"),
		"creation": ("Date and time the record was created.", moment, "Datetime"),
		"modified": ("Date and time the record was last changed.", moment, "Datetime"),
		"modified_by": ("User who last changed the record.", "Administrator", "Data"),
		"idx": ("Position of the record within its list or table.", 1, "Int"),
		"parent": (
			"Name of the document this row belongs to.",
			str(context.get("parent_name", "Example")),
			"Data",
		),
		"parentfield": (
			"Field of the parent document that holds this table.",
			context.get("fieldname", "Example"),
			"Data",
		),
		"parenttype": (
			"DocType of the document this row belongs to.",
			context.get("doctype", "Example"),
			"Data",
		),
	}
	submittable = bool(cint(meta.is_submittable))
	info["docstatus"] = (
		"Document status: 0 = Draft, 1 = Submitted, 2 = Cancelled."
		if submittable
		else "Document status. Always 0, because this DocType cannot be submitted.",
		0,
		"Int",
	)

	records = []
	for name in STANDARD_FIELDS_CHILD if is_child else STANDARD_FIELDS_PARENT:
		if name not in column_types:
			continue
		text, example, fieldtype = info[name]
		record = {
			"fieldname": name,
			"df": frappe._dict(fieldname=name, label=name, fieldtype=fieldtype, options=None),
			"db_type": column_types[name],
			"required": False,
			"conditional": False,
			"is_table": False,
			"description": f"{text} Set by the system. Read only.",
			"example": example,
		}
		if name == "docstatus":
			record["schema_extra"] = {"enum": [0, 1, 2] if submittable else [0]}
		records.append(record)
	return records


def database_version():
	"""The database engine's version number, e.g. 10.11.14, or "unknown" if it cannot be read."""
	query = "select sqlite_version()" if frappe.db.db_type == "sqlite" else "select version()"
	try:
		found = re.match(r"\d+(\.\d+)*", str(frappe.db.sql(query)[0][0]))
	except Exception:
		found = None
	return found.group(0) if found else "unknown"


def generation_stamp(include_standard=False, customisations_only=False):
	"""Where and when the schema was generated, so a reader knows which versions it describes.

	Deliberately has no site name, company or user: it must be safe to hand to a third party.
	"""
	apps = {}
	for app in frappe.get_installed_apps():
		try:
			apps[app] = frappe.get_module(app).__version__
		except Exception:
			apps[app] = "unknown"
	return {
		"date": today(),
		"frappe": apps.get("frappe"),
		"apps": apps,
		"database": f"{frappe.db.db_type} {database_version()}",
		"includes_standard_fields": bool(include_standard),
		"customisations_only": bool(customisations_only),
	}


def property_setter_fields(doctype):
	"""Fields of a DocType whose properties a Property Setter changes (e.g. a hidden or relabelled field)."""
	return frappe.get_all(
		"Property Setter", filters={"doc_type": doctype, "field_name": ("is", "set")}, pluck="field_name"
	)


def custom_sources(doctype):
	"""{fieldname: who added it} for the DocType's custom fields.

	Frappe records the module (so the app) only sometimes; it always records whether code created
	the field (an installed app) or a person did, so that is the fallback. We never guess an app name.
	"""
	sources = {}
	for row in frappe.get_all(
		"Custom Field", filters={"dt": doctype}, fields=["fieldname", "module", "is_system_generated"]
	):
		app = frappe.db.get_value("Module Def", row.module, "app_name") if row.module else None
		sources[row.fieldname] = (
			f"added by {app}" if app else "added by an app" if row.is_system_generated else "added by a user"
		)
	return sources


def humanize_condition(expression):
	"""Turn a mandatory_depends_on expression into a sentence, or say only that it is conditional.

	Raw expressions are code, so anything we cannot read reliably is not shown.
	"""
	expr = (expression or "").strip().rstrip(";").strip()
	if not expr:
		return None
	if re.fullmatch(r"[A-Za-z_]\w*", expr):
		return f"Required when '{expr}' is set."
	body = expr[5:].strip() if expr.startswith("eval:") else expr
	if match := re.fullmatch(r"(!?)\s*doc\.(\w+)", body):
		return f"Required when '{match.group(2)}' is {'not set' if match.group(1) else 'set'}."
	if match := re.fullmatch(r"doc\.(\w+)\s*(==|!=)\s*(?:'([^']*)'|\"([^\"]*)\"|(\d+))", body):
		field, operator = match.group(1), match.group(2)
		value = next(v for v in match.groups()[2:] if v is not None)
		return f"Required when '{field}' {'is' if operator == '==' else 'is not'} '{value}'."
	return "Required under some conditions."


def sentence(text):
	"""Plain text with one trailing full stop."""
	text = " ".join(strip_html_tags(text or "").split())
	return text if not text or text[-1] in ".!?" else text + "."


def describe(df, custom_source, doctype):
	"""One cell of text saying what a field is. `custom_source` is None for a standard field."""
	if df.fieldtype in table_fields:
		text = (
			f"Child table linking rows of doctype '{df.options}' to this {doctype}; "
			f"see fields below suffixed (child: {df.options})."
		)
	else:
		label = df.label or df.fieldname
		text = (
			label if label[-1:] in ".?!:" else f"{label}."
		)  # a label already ending in "?" needs no full stop
		if df.fieldtype == "Link" and df.options:
			text += f" Links to the '{df.options}' doctype."
		elif df.fieldtype == "Dynamic Link" and df.options:
			text += f" Links to the doctype named in '{df.options}'."
		elif df.fieldtype == "Select":
			options = select_options(df)
			if options:
				text += f" One of: {', '.join(options)}."
		if help_text := sentence(df.description):
			text += f" {help_text}"
		if cint(df.read_only):
			text += " Read only."
		if cint(df.set_only_once):
			text += " Can only be set once."
		if cint(df.unique):
			text += " Must be unique."
		if df.fetch_from:
			text += f" Fetched from {df.fetch_from}."
		if df.default not in (None, ""):
			text += f" Default: {DYNAMIC_DEFAULTS.get(df.default, df.default)}."
		if condition := humanize_condition(df.mandatory_depends_on):
			text += f" {condition}"
	if custom_source is None:
		return text
	return f"(custom field, {custom_source}) {text}"


def select_options(df):
	return [o.strip() for o in (df.options or "").split("\n") if o.strip()]


def name_description(meta):
	autoname = (meta.autoname or "").strip()
	base = f"Frappe's autogenerated primary key for {meta.name}."
	if autoname.startswith("naming_series:"):
		return f"{base} Autonamed via naming_series pattern '{naming_series_options(meta)[0] if naming_series_options(meta) else ''}'."
	if autoname.startswith("field:"):
		return f"{base} Taken from the '{autoname[6:]}' field."
	if autoname.startswith("format:"):
		return f"{base} Autonamed with the format '{autoname[7:]}'."
	if autoname == "prompt":
		return f"{base} Entered by the user when the record is created."
	if autoname == "hash":
		return f"{base} Autogenerated random hash."
	return base  # named by controller code or by a rule we cannot see, so claim nothing more


def naming_series_options(meta):
	df = meta.get_field("naming_series")
	return select_options(df) if df else []


# ---- example values ---------------------------------------------------------------------


def name_example(meta, db_type):
	"""An example `name` of the right JSON type: auto-increment DocTypes have an integer name."""
	if json_type(db_type) == "integer":
		return 1
	return example_name(meta)


def example_name(meta):
	autoname = (meta.autoname or "").strip()
	year = today()[:4]
	if autoname.startswith("naming_series:") and naming_series_options(meta):
		return render_series(naming_series_options(meta)[0], year, add_counter=True)
	if autoname.startswith("format:"):
		return render_series(autoname[7:], year)
	if autoname.startswith("field:"):
		return f"Example {autoname[6:].replace('_', ' ')}"
	return f"{meta.name.replace(' ', '-').upper()}-{year}-00001"


def render_series(pattern, year, add_counter=False):
	"""Turn a Frappe naming pattern like ACC-SINV-.YYYY.-.##### into a sample name.

	A naming_series option without #'s gets Frappe's default 5-digit counter appended.
	"""
	out = pattern.replace(".YYYY.", year).replace("{YYYY}", year).replace(".YY.", year[2:])
	if add_counter and "#" not in out:
		out += "00001"
	out = re.sub(r"\.?(#+)\.?", lambda m: "1".zfill(len(m.group(1))), out)
	out = re.sub(r"\{#+\}", lambda m: "1".zfill(len(m.group(0)) - 2), out)
	return out.replace("..", ".").strip(".")


def default_masters():
	"""Stable master records that make link examples believable, taken from the site's own settings."""
	# Global Defaults is an ERPNext DocType; a site without ERPNext must still be able to export
	currency = None
	if frappe.db.exists("DocType", "Global Defaults"):
		currency = frappe.db.get_single_value("Global Defaults", "default_currency")
	return {
		"Currency": currency or "USD",
		"Country": frappe.db.get_single_value("System Settings", "country") or "United States",
		"Language": frappe.db.get_single_value("System Settings", "language") or "en",
		"UOM": "Nos",  # ERPNext's standard unit
		"Company": "Example Company",
	}


def example_value(df, db_type, masters=None):
	"""A typed, synthetic example for a field. Never real data, never null.

	Amounts, quantities and dates are chosen so a whole sample record reads sensibly: nothing is
	paid yet, quantity times rate equals the total, and due dates fall after posting dates.
	"""
	masters = masters or {}
	target = json_type(db_type, df.fieldtype)
	label = df.label or df.fieldname
	ftype = df.fieldtype
	options = df.options or ""

	if df.default and df.default not in ("Today", "Now", "__user", "user"):
		cast = cast_to(df.default, target)
		if cast is not None:
			return cast

	if ftype == "Select":
		choices = select_options(df)
		return choices[0] if choices else f"Example {label}"
	if ftype == "Link":
		if not options:
			return "Example"
		if options in masters:
			return masters[options]
		if options == "Price List":
			return "Standard Buying" if "buying" in df.fieldname else "Standard Selling"
		return f"Example {options}"
	if ftype == "Dynamic Link":
		return "Example"
	if ftype == "Date":
		return add_days(today(), 30) if LATER_DATE.search(df.fieldname) else today()
	if ftype == "Datetime":
		day = add_days(today(), 30) if LATER_DATE.search(df.fieldname) else today()
		return f"{day} 10:00:00"
	if ftype == "Time":
		return "10:00:00"
	if ftype == "Check":
		return 0
	if ftype in ("Int", "Long Int"):
		return 1
	if ftype in ("Currency", "Float", "Percent") and (
		ZERO_AMOUNT.search(df.fieldname) or PROGRESS.search(df.fieldname)
	):
		return 0
	if ftype == "Float" and (RATE_OF_ONE.search(df.fieldname) or QUANTITY.search(df.fieldname)):
		return 1
	if ftype in ("Currency", "Float"):
		return 1000 if ftype == "Currency" else 1
	if ftype == "Percent":
		return 10
	if ftype == "Rating":
		return 0.5
	if ftype == "Duration":
		return 3600
	if ftype in ("Attach", "Attach Image"):
		return "/files/example.pdf"
	if ftype in ("Code", "JSON"):
		return "{}"
	if ftype == "Data" and options == "Email":
		return "user@example.com"
	if ftype == "Data" and options == "Phone":
		return "+1-555-0100"
	if ftype == "Data" and options == "URL":
		return "https://example.com"
	if target == "number":
		return 1
	if target == "integer":
		return 1
	return f"Example {label}"


def cast_to(value, target):
	"""Convert to the JSON type, or None if it does not fit."""
	try:
		if target == "integer":
			return int(float(value))
		if target == "number":
			number = float(value)
			return int(number) if number == int(number) else number
	except (TypeError, ValueError):
		return None
	return str(value)


# ---- payload and schema -----------------------------------------------------------------


def find_doctype(name):
	"""The DocType's real name for a typed name, ignoring capitalisation on every database engine."""
	actual = frappe.db.get_value("DocType", name, "name")
	if actual:
		return actual
	# MariaDB already matched regardless of case above; SQLite and Postgres compare exactly
	lowered = name.lower()
	return next((n for n in frappe.get_all("DocType", pluck="name") if n.lower() == lowered), None)


def resolve_doctype(name):
	"""The DocType's real name for whatever was typed: trims spaces and fixes capitalisation."""
	name = (name or "").strip()
	if not name:
		frappe.throw(_("Select a DocType first"))
	actual = find_doctype(name)
	if not actual:
		frappe.throw(_("DocType {0} does not exist").format(frappe.bold(name)))
	return actual


def build_document(doctype, include_standard=False, customisations_only=False):
	"""Everything the three sheets are rendered from, for any DocType."""
	frappe.only_for("System Manager", message=True)
	doctype = resolve_doctype(doctype)

	fields = collect_fields(
		doctype,
		is_child=bool(cint(frappe.get_meta(doctype).istable)),
		include_standard=include_standard,
		customisations_only=customisations_only,
	)
	parent_name = next((r["example"] for r in fields if r["fieldname"] == "name"), "Example")
	children = {}
	for record in fields:
		if record["is_table"] and record["child_doctype"] not in children:
			child = record["child_doctype"]
			# if two tables share a child DocType, its rows describe the first of them
			context = {"doctype": doctype, "fieldname": record["fieldname"], "parent_name": parent_name}
			children[child] = collect_fields(child, True, include_standard, context, customisations_only)

	if customisations_only:
		# a table stays if it was itself customised or its child DocType has something customised
		system = set(STANDARD_FIELDS_PARENT) | set(STANDARD_FIELDS_CHILD)
		fields = [
			f
			for f in fields
			if not f["is_table"]
			or f["customised"]
			or any(r["fieldname"] not in system for r in children[f["child_doctype"]])
		]
		kept = {f["child_doctype"] for f in fields if f["is_table"]}
		children = {name: rows for name, rows in children.items() if name in kept}

	return {
		"doctype": doctype,
		"fields": fields,
		"children": children,
		"include_standard": bool(include_standard),
		"customisations_only": bool(customisations_only),
	}


def build_payload(document):
	def row_from(records):
		return {r["fieldname"]: r["example"] for r in records if not r["is_table"]}

	record = {}
	for field in document["fields"]:
		if field["is_table"]:
			record[field["fieldname"]] = [row_from(document["children"][field["child_doctype"]])]
		else:
			record[field["fieldname"]] = field["example"]
	return {document["doctype"]: [record]}


def required_label(record):
	"""Yes, or Conditional when only some situations make it mandatory, or No."""
	return "Yes" if record["required"] else "Conditional" if record.get("conditional") else "No"


def select_values(df):
	"""The values a Select field accepts. An empty value counts only if the options start with a blank line."""
	raw = [o.strip() for o in (df.options or "").split("\n")]
	values = [o for o in raw if o]
	if raw and raw[0] == "" and values:
		values.insert(0, "")
	return list(dict.fromkeys(values))


def schema_for(record):
	"""JSON Schema for one scalar field: type, plus the constraints the metadata really states."""
	db_type, df = record["db_type"], record["df"]
	schema = {
		"type": json_type(db_type, df.fieldtype),
		"description": f"{db_type}. Required: {required_label(record)}. {record['description']}",
	}
	if df.fieldtype == "Select" and (values := select_values(df)):
		schema["enum"] = values
	elif df.fieldtype == "Date":
		schema["format"] = "date"
	elif df.fieldtype == "Data" and df.options == "Email":
		schema["format"] = "email"
	elif df.fieldtype == "Data" and df.options == "URL":
		schema["format"] = "uri"
	if schema["type"] == "string" and (length := re.fullmatch(r"varchar\((\d+)\)", db_type)):
		schema["maxLength"] = int(length.group(1))
	schema.update(record.get("schema_extra", {}))
	return schema


def build_schema(document):
	doctype = document["doctype"]

	def object_schema(records):
		properties = {}
		for r in records:
			if r["is_table"]:
				child_records = document["children"][r["child_doctype"]]
				child = object_schema(child_records)
				properties[r["fieldname"]] = {
					"type": "array",
					"description": f"Child table -> doctype '{r['child_doctype']}'.",
					"items": {"type": "object", **child},
				}
			else:
				properties[r["fieldname"]] = schema_for(r)
		return {"properties": properties, "required": [r["fieldname"] for r in records if r["required"]]}

	tables = [f["fieldname"] for f in document["fields"] if f["is_table"]]
	item = {
		"title": doctype,
		"description": f"ERPNext {doctype} document"
		+ (f" with child tables ({', '.join(tables)})." if tables else "."),
		"type": "object",
		**object_schema(document["fields"]),
	}
	stamp = generation_stamp(document.get("include_standard"), document.get("customisations_only"))
	versions = ", ".join(f"{app} {version}" for app, version in stamp["apps"].items())
	return {
		"$schema": "https://json-schema.org/draft/2020-12/schema",
		"title": f"{doctype} Payload",
		"description": (
			f"Schema of a {doctype} document, generated {stamp['date']} on {stamp['database']} "
			f"with {versions}. Data types are as reported by this database engine."
			+ (
				" Only fields added or changed by customisation are listed."
				if stamp["customisations_only"]
				else ""
			)
		),
		"x-generated-from": stamp,
		"type": "object",
		"properties": {doctype: {"type": "array", "items": item}},
		"required": [doctype],
	}


# ---- workbook ---------------------------------------------------------------------------


def sheet_title(doctype):
	"""Excel sheet names: at most 31 characters, none of []:*?/\\"""
	return INVALID_SHEET_CHARS.sub(" ", doctype).strip()[:31] or "Schema"


def file_stem(doctype):
	return re.sub(r"\W+", "_", doctype).strip("_") + "_Global_Schema"


def field_rows(document):
	rows = []
	for record in document["fields"]:
		rows.append(
			(
				record["fieldname"],
				record["db_type"],
				required_label(record),
				record["description"],
				None if record["is_table"] else record["example"],
			)
		)
		if record["is_table"]:
			child = record["child_doctype"]
			for child_record in document["children"][child]:
				rows.append(
					(
						f"{child_record['fieldname']} (child: {child})",
						child_record["db_type"],
						required_label(child_record),
						child_record["description"],
						None if child_record["is_table"] else child_record["example"],
					)
				)
	return rows


def safe_cell(value):
	"""Strip characters Excel rejects and cap the length of a cell."""
	if isinstance(value, str):
		value = ILLEGAL_CHARACTERS_RE.sub("", value)
		if len(value) > MAX_CELL_CHARS:
			value = value[: MAX_CELL_CHARS - 1] + "…"
	return value


def append_row(sheet, values):
	"""Append a row of safe cells. Text that starts with "=" is stored as text, never as a formula.

	Labels, defaults and Select options come from DocType definitions, and the workbook is meant to
	be shared, so a definition must not be able to put a formula into the reader's Excel.
	"""
	sheet.append([safe_cell(value) for value in values])
	for cell in sheet[sheet.max_row]:
		if isinstance(cell.value, str) and cell.value.startswith("="):
			cell.data_type = "s"


def write_json_sheet(worksheet, data):
	for line in json.dumps(data, indent=2, ensure_ascii=False).split("\n"):
		append_row(worksheet, [line])
	worksheet.column_dimensions["A"].width = TEXT_SHEET_WIDTH


def build_workbook(doctype, include_standard=False, customisations_only=False):
	"""Return (xlsx bytes, filename) for a DocType."""
	return render_workbook(build_document(doctype, include_standard, customisations_only))


def render_workbook(document):
	"""Turn a built document into (xlsx bytes, filename)."""
	doctype = document["doctype"]

	workbook = Workbook()
	sheet = workbook.active
	sheet.title = sheet_title(doctype)
	sheet.append(HEADERS)
	bold, yellow = Font(bold=True), PatternFill("solid", fgColor="FFFFFF00")
	for cell in sheet[1]:
		cell.font, cell.fill = bold, yellow
	for row in field_rows(document):
		append_row(sheet, row)
	for letter, width in zip("ABCDE", COLUMN_WIDTHS, strict=False):
		sheet.column_dimensions[letter].width = width

	write_json_sheet(workbook.create_sheet("Payload"), build_payload(document))
	write_json_sheet(workbook.create_sheet("Schema"), build_schema(document))

	buffer = io.BytesIO()
	workbook.save(buffer)
	return buffer.getvalue(), file_stem(doctype) + ".xlsx"


def summarize(document):
	"""Counts for the preview, taken from the very rows that go into the file."""
	meta = frappe.get_meta(document["doctype"])
	scalars = [f for f in document["fields"] if not f["is_table"]]
	kind = (
		"Single DocType"
		if cint(meta.issingle)
		else "Child table"
		if cint(meta.istable)
		else "Virtual DocType"
		if cint(meta.get("is_virtual"))
		else "DocType"
	)
	rows = len(field_rows(document))
	return {
		"doctype": document["doctype"],
		"module": meta.module,
		"kind": kind,
		"custom": bool(cint(meta.custom)),
		"submittable": bool(cint(meta.is_submittable)),
		"fields": len(scalars),
		"required": sum(1 for f in scalars if f["required"]),
		"conditional": sum(1 for f in scalars if f.get("conditional")),
		"custom_fields": sum(1 for f in scalars if f["description"].startswith("(custom field")),
		"tables": [
			{
				"fieldname": f["fieldname"],
				"doctype": f["child_doctype"],
				"fields": sum(1 for c in document["children"][f["child_doctype"]] if not c["is_table"]),
			}
			for f in document["fields"]
			if f["is_table"]
		],
		"rows": rows,
		"empty": rows == 0,
		"include_standard": document["include_standard"],
		"customisations_only": document["customisations_only"],
	}


@frappe.whitelist()
def get_preview(doctype: str, include_standard: int = 0, customisations_only: int = 0):
	"""What the workbook for this DocType would contain, without building the file."""
	return summarize(build_document(doctype, cint(include_standard), cint(customisations_only)))


def send_file(content, filename):
	frappe.response["filename"] = filename
	frappe.response["filecontent"] = content
	frappe.response["type"] = "binary"


@frappe.whitelist()
def download_workbook(doctype: str, include_standard: int = 0, customisations_only: int = 0):
	"""Send one workbook to the browser as a file download. Works for any DocType."""
	send_file(*build_workbook(doctype, cint(include_standard), cint(customisations_only)))


# ---- several DocTypes at once -----------------------------------------------------------


def build_pack(doctypes, include_standard=False, customisations_only=False):
	"""A zip with one workbook per DocType plus an index workbook. Returns (bytes, filename)."""
	frappe.only_for("System Manager", message=True)

	wanted = list(dict.fromkeys(str(d).strip() for d in (doctypes or []) if str(d).strip()))
	if not wanted:
		frappe.throw(_("Select at least one DocType"))
	if len(wanted) > MAX_PACK:
		frappe.throw(_("Export at most {0} DocTypes at a time").format(MAX_PACK))

	resolved, unknown = [], []
	for name in wanted:
		actual = find_doctype(name)
		(resolved if actual else unknown).append(actual or name)
	if unknown:
		frappe.throw(_("These DocTypes do not exist: {0}").format(", ".join(frappe.bold(n) for n in unknown)))
	resolved = list(dict.fromkeys(resolved))  # two spellings of one DocType are one DocType

	archive_buffer = io.BytesIO()
	index = [("DocType", "Module", "Kind", "Fields", "Required", "Child tables", "Rows", "File")]
	used = set()
	with zipfile.ZipFile(archive_buffer, "w", zipfile.ZIP_DEFLATED) as archive:
		for name in resolved:
			document = build_document(name, include_standard, customisations_only)
			content, filename = render_workbook(document)
			while filename in used:  # two DocTypes can sanitise to the same file name
				filename = filename.replace(".xlsx", "_2.xlsx")
			used.add(filename)
			archive.writestr(filename, content)
			info = summarize(document)
			index.append(
				(
					info["doctype"],
					info["module"],
					info["kind"],
					info["fields"],
					info["required"],
					len(info["tables"]),
					info["rows"],
					filename,
				)
			)
		archive.writestr("_Index.xlsx", render_index(index))

	return archive_buffer.getvalue(), f"Fieldbook_{len(resolved)}_DocTypes_{today().replace('-', '')}.zip"


def render_index(rows):
	workbook = Workbook()
	sheet = workbook.active
	sheet.title = "Index"
	bold, yellow = Font(bold=True), PatternFill("solid", fgColor="FFFFFF00")
	for row in rows:
		append_row(sheet, row)
	for cell in sheet[1]:
		cell.font, cell.fill = bold, yellow
	for letter, width in zip("ABCDEFGH", (32, 18, 16, 8, 10, 13, 8, 46), strict=False):
		sheet.column_dimensions[letter].width = width
	buffer = io.BytesIO()
	workbook.save(buffer)
	return buffer.getvalue()


@frappe.whitelist()
def download_pack(doctypes: str, include_standard: int = 0, customisations_only: int = 0):
	"""Send a zip of workbooks for several DocTypes."""
	frappe.only_for("System Manager", message=True)
	try:
		names = frappe.parse_json(doctypes)
	except ValueError:  # not JSON at all
		names = None
	if not isinstance(names, list):
		frappe.throw(_("Select at least one DocType"))
	send_file(*build_pack(names, cint(include_standard), cint(customisations_only)))
