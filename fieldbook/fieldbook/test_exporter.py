# Copyright (c) 2026, Sidharth PV and contributors
# See license.txt

import functools
import io
import json
import re
import zipfile
from unittest.mock import patch

import frappe
from frappe.tests.utils import FrappeTestCase
from openpyxl import load_workbook

from fieldbook.fieldbook.exporter import (
	CHILD_TABLE_TYPE,
	MAX_PACK,
	SKIP_FIELDTYPES,
	STANDARD_FIELDS_CHILD,
	STANDARD_FIELDS_PARENT,
	build_document,
	build_pack,
	build_payload,
	build_schema,
	build_workbook,
	collect_fields,
	custom_sources,
	database_version,
	default_masters,
	describe,
	download_pack,
	download_workbook,
	example_value,
	file_stem,
	find_doctype,
	generation_stamp,
	get_column_types,
	get_preview,
	humanize_condition,
	naming_series_options,
	property_setter_fields,
	render_series,
	required_label,
	resolve_doctype,
	safe_cell,
	select_values,
	sheet_title,
	type_from_fieldtype,
)

PY_TYPES = {"string": str, "integer": int, "number": (int, float), "array": list, "object": dict}


def validate(value, schema, path="$"):
	"""Minimal JSON Schema check (type, properties, required, items), enough for our own output."""
	errors = []
	expected = PY_TYPES[schema["type"]]
	if isinstance(value, bool) or not isinstance(value, expected):
		return [f"{path}: expected {schema['type']}, got {type(value).__name__}"]
	if "enum" in schema and value not in schema["enum"]:
		errors.append(f"{path}: {value!r} not in enum")
	if "maxLength" in schema and len(value) > schema["maxLength"]:
		errors.append(f"{path}: longer than maxLength {schema['maxLength']}")
	if schema.get("format") == "date" and not re.fullmatch(r"\d{4}-\d\d-\d\d", value):
		errors.append(f"{path}: {value!r} is not a date")
	if schema["type"] == "object":
		for key in schema.get("required", []):
			if key not in value:
				errors.append(f"{path}: missing required {key}")
		for key, sub in value.items():
			if key not in schema["properties"]:
				errors.append(f"{path}: unexpected {key}")
			else:
				errors += validate(sub, schema["properties"][key], f"{path}.{key}")
	if schema["type"] == "array":
		for i, item in enumerate(value):
			errors += validate(item, schema["items"], f"{path}[{i}]")
	return errors


def needs_erpnext(test):
	"""Skip on a site without ERPNext, where DocTypes such as Sales Invoice do not exist."""

	@functools.wraps(test)
	def wrapper(self, *args, **kwargs):
		if not frappe.db.exists("DocType", "Sales Invoice"):
			self.skipTest("needs ERPNext")
		return test(self, *args, **kwargs)

	return wrapper


class TestExporter(FrappeTestCase):
	def tearDown(self):
		frappe.set_user("Administrator")

	def workbook(self, doctype):
		content, filename = build_workbook(doctype)
		return load_workbook(io.BytesIO(content)), filename

	def workbook_with_standard(self, doctype):
		content, filename = build_workbook(doctype, include_standard=True)
		return load_workbook(io.BytesIO(content)), filename

	def assert_valid(self, doctype):
		document = build_document(doctype)
		self.assertEqual(validate(build_payload(document), build_schema(document)), [], doctype)
		return document

	# workbook layout

	@needs_erpnext
	def test_workbook_has_three_sheets_headers_and_widths(self):
		wb, filename = self.workbook("Sales Invoice")
		self.assertEqual(wb.sheetnames, ["Sales Invoice", "Payload", "Schema"])
		self.assertEqual(filename, "Sales_Invoice_Global_Schema.xlsx")
		sheet = wb["Sales Invoice"]
		self.assertEqual(
			[c.value for c in sheet[1]], ["Field Name", "Data Type", "Required", "Description", "Example"]
		)
		self.assertTrue(all(c.font.b and c.fill.fgColor.rgb == "FFFFFF00" for c in sheet[1]))
		self.assertEqual([sheet.column_dimensions[c].width for c in "ABCDE"], [34, 16, 10, 95, 34])
		self.assertEqual(wb["Payload"].column_dimensions["A"].width, 120)
		self.assertEqual(wb["Schema"].column_dimensions["A"].width, 120)

	@needs_erpnext
	def test_name_row_comes_first_and_required_fields_precede_optional_ones(self):
		sheet = self.workbook("Sales Invoice")[0]["Sales Invoice"]
		rows = [[c.value for c in r] for r in sheet.iter_rows(min_row=2)]
		self.assertEqual(rows[0][0], "name")
		parent = [r for r in rows if "(child:" not in r[0] and r[1] != CHILD_TABLE_TYPE]
		flags = [r[2] == "Yes" for r in parent]
		self.assertEqual(flags, sorted(flags, reverse=True))  # every required field before any other
		self.assertLessEqual({r[2] for r in rows}, {"Yes", "No", "Conditional"})

	@needs_erpnext
	def test_data_type_is_the_real_database_column_type(self):
		sheet = self.workbook("Sales Invoice")[0]["Sales Invoice"]
		types = get_column_types("Sales Invoice")
		checked = 0
		for row in sheet.iter_rows(min_row=2, values_only=True):
			if "(child:" in row[0] or row[1] == CHILD_TABLE_TYPE:
				continue
			self.assertEqual(row[1], types[row[0]], row[0])
			checked += 1
		self.assertGreater(checked, 50)

	@needs_erpnext
	def test_child_tables_are_flattened_after_their_table_row(self):
		sheet = self.workbook("Sales Invoice")[0]["Sales Invoice"]
		names = [r[0] for r in sheet.iter_rows(min_row=2, values_only=True)]
		items_at = names.index("items")
		row = [c.value for c in sheet[items_at + 2]]
		self.assertEqual(row[1], CHILD_TABLE_TYPE)
		self.assertIsNone(row[4])  # a table row has no scalar example
		self.assertTrue(names[items_at + 1].endswith("(child: Sales Invoice Item)"))
		self.assertIn("qty (child: Sales Invoice Item)", names[items_at:])

	@needs_erpnext
	def test_hidden_layout_and_virtual_fields_are_excluded(self):
		meta = frappe.get_meta("Sales Invoice")
		hidden = [df.fieldname for df in meta.fields if df.hidden and df.fieldtype not in SKIP_FIELDTYPES]
		self.assertTrue(hidden)
		exported = {r["fieldname"] for r in collect_fields("Sales Invoice")}
		self.assertFalse(exported & set(hidden))
		layout = {df.fieldname for df in meta.fields if df.fieldtype in SKIP_FIELDTYPES}
		self.assertFalse(exported & layout)

	@needs_erpnext
	def test_readonly_fields_are_included(self):
		self.assertIn("base_grand_total", {r["fieldname"] for r in collect_fields("Sales Invoice")})

	def test_sheet_title_and_filename_are_safe(self):
		self.assertEqual(len(sheet_title("Sales Taxes and Charges Template Extra Long")), 31)
		self.assertNotIn("/", sheet_title("A/B:C*D?E[F]"))
		self.assertEqual(file_stem("Sales Taxes/and Charges"), "Sales_Taxes_and_Charges_Global_Schema")

	# descriptions

	@needs_erpnext
	def test_descriptions_include_link_and_select_hints(self):
		by_name = {r["fieldname"]: r for r in collect_fields("Sales Invoice")}
		self.assertEqual(by_name["company"]["description"], "Company. Links to the 'Company' doctype.")
		self.assertIn("One of:", by_name["naming_series"]["description"])
		# the pattern is whatever this site declares first (ACC-SINV-.YYYY.- on v15, SINV-.YY.- on v16)
		series = naming_series_options(frappe.get_meta("Sales Invoice"))[0]
		self.assertIn(f"naming_series pattern '{series}'", by_name["name"]["description"])

	def test_name_description_only_claims_a_pattern_the_doctype_declares(self):
		by_name = lambda dt: {r["fieldname"]: r for r in collect_fields(dt)}["name"]["description"]  # noqa: E731
		self.assertEqual(by_name("Contact"), "Frappe's autogenerated primary key for Contact.")
		self.assertIn("random hash", by_name("ToDo"))

	def test_a_label_that_already_ends_in_punctuation_gets_no_extra_full_stop(self):
		def df(label):
			return frappe._dict(
				fieldname="f", label=label, fieldtype="Data", options=None, mandatory_depends_on=None
			)

		self.assertEqual(describe(df("Is exempt?"), None, "X"), "Is exempt?")
		self.assertEqual(describe(df("Rate"), None, "X"), "Rate.")
		self.assertEqual(describe(df("Name:"), None, "X"), "Name:")

	def test_custom_field_prefix(self):
		df = frappe._dict(fieldname="x", label="X", fieldtype="Data", options=None, mandatory_depends_on=None)
		self.assertTrue(
			describe(df, "added by an app", "ToDo").startswith("(custom field, added by an app) X.")
		)
		self.assertFalse(describe(df, None, "ToDo").startswith("(custom"))

	def test_custom_sources_never_guess_an_app_name(self):
		rows = [
			frappe._dict(fieldname="a", module="Accounts", is_system_generated=1),
			frappe._dict(fieldname="b", module=None, is_system_generated=1),
			frappe._dict(fieldname="c", module=None, is_system_generated=0),
		]
		with patch("frappe.get_all", return_value=rows), patch("frappe.db.get_value", return_value="erpnext"):
			sources = custom_sources("Anything")
		self.assertEqual(sources, {"a": "added by erpnext", "b": "added by an app", "c": "added by a user"})

	@needs_erpnext
	def test_real_custom_fields_say_who_added_them(self):
		fields = {r["fieldname"]: r for r in collect_fields("Sales Invoice")}
		if "exempt_from_sales_tax" not in fields:
			self.skipTest("ERPNext custom field not present")
		self.assertTrue(fields["exempt_from_sales_tax"]["description"].startswith("(custom field, added by"))

	# examples

	def test_synthetic_examples_by_type(self):
		def df(fieldtype, options=None, default=None):
			return frappe._dict(
				fieldname="f", label="F", fieldtype=fieldtype, options=options, default=default
			)

		self.assertEqual(example_value(df("Select", "A\nB"), "varchar(140)"), "A")
		self.assertEqual(example_value(df("Link", "Customer"), "varchar(140)"), "Example Customer")
		self.assertEqual(example_value(df("Link", "UOM"), "varchar(140)", {"UOM": "Nos"}), "Nos")
		self.assertEqual(example_value(df("Date"), "date"), frappe.utils.today())
		self.assertEqual(example_value(df("Check"), "int(1)"), 0)
		self.assertEqual(example_value(df("Data", "Email"), "varchar(140)"), "user@example.com")
		self.assertEqual(example_value(df("Data"), "varchar(140)"), "Example F")

	def test_render_series(self):
		# a bare series gets Frappe's default 5-digit counter, exactly as in real names
		self.assertEqual(render_series("ACC-SINV-.YYYY.-", "2026", add_counter=True), "ACC-SINV-2026-00001")
		self.assertEqual(render_series("ACC-SINV-.YYYY.-", "2026"), "ACC-SINV-2026-")
		self.assertEqual(render_series("SO-.#####", "2026"), "SO-00001")

	# payload and schema

	@needs_erpnext
	def test_payload_has_no_nulls_and_nests_child_tables(self):
		record = build_payload(build_document("Sales Invoice"))["Sales Invoice"][0]
		self.assertIsInstance(record["items"], list)
		self.assertEqual(len(record["items"]), 1)
		self.assertNotIn(None, record.values())
		self.assertNotIn(None, record["items"][0].values())

	@needs_erpnext
	def test_payload_validates_against_schema(self):
		for doctype in ("Sales Invoice", "ToDo", "Contact"):
			self.assert_valid(doctype)

	def test_validator_actually_catches_problems(self):
		document = build_document("ToDo")
		schema, payload = build_schema(document), build_payload(document)
		payload["ToDo"][0].pop("name")
		self.assertTrue(validate(payload, schema))

	@needs_erpnext
	def test_schema_shape(self):
		schema = build_schema(build_document("Sales Invoice"))
		self.assertEqual(schema["$schema"], "https://json-schema.org/draft/2020-12/schema")
		self.assertEqual(schema["title"], "Sales Invoice Payload")
		self.assertEqual(schema["required"], ["Sales Invoice"])
		item = schema["properties"]["Sales Invoice"]["items"]
		self.assertIn("name", item["required"])
		child = item["properties"]["items"]
		self.assertEqual(child["type"], "array")
		self.assertIn("item_name", child["items"]["required"])
		self.assertIn("Required: Yes.", item["properties"]["company"]["description"])

	def test_json_sheets_round_trip(self):
		wb = self.workbook("ToDo")[0]
		for name in ("Payload", "Schema"):
			text = "\n".join(str(r[0]) for r in wb[name].iter_rows(values_only=True))
			self.assertIsInstance(json.loads(text), dict)

	@needs_erpnext
	def test_other_database_engines_fall_back_to_frappes_own_table_description(self):
		described = [
			{"name": "grand_total", "type": "numeric(21,9)"},
			{"name": "name", "type": "varchar(140)"},
		]
		with (
			patch.object(frappe.db, "db_type", "postgres"),
			patch.object(frappe.db, "get_table_columns_description", return_value=described) as helper,
		):
			self.assertEqual(
				get_column_types("Sales Invoice"), {"grand_total": "numeric(21,9)", "name": "varchar(140)"}
			)
		helper.assert_called_once_with("tabSales Invoice")
		# and an engine that cannot describe the table degrades to "no table" instead of crashing
		with (
			patch.object(frappe.db, "db_type", "postgres"),
			patch.object(frappe.db, "get_table_columns_description", side_effect=Exception("boom")),
		):
			self.assertEqual(get_column_types("Sales Invoice"), {})

	def test_numeric_and_real_columns_are_json_numbers(self):
		from fieldbook.fieldbook.exporter import json_type

		for db_type in ("decimal(21,9)", "numeric(21,9)", "double precision", "real", "float"):
			self.assertEqual(json_type(db_type), "number", db_type)
		for db_type in ("int(11)", "bigint", "smallint", "integer", "tinyint(4)"):
			self.assertEqual(json_type(db_type), "integer", db_type)
		self.assertEqual(json_type("varchar(140)"), "string")

	def test_names_resolve_regardless_of_case_on_engines_that_compare_exactly(self):
		real = frappe.db.get_value

		def exact_only(doctype, name, *args, **kwargs):
			return real(doctype, name, *args, **kwargs) if name in ("ToDo", "Contact") else None

		with patch("frappe.db.get_value", side_effect=exact_only):
			self.assertEqual(find_doctype("todo"), "ToDo")
			self.assertEqual(find_doctype("CONTACT"), "Contact")
			self.assertEqual(find_doctype("ToDo"), "ToDo")
			self.assertIsNone(find_doctype("no such doctype"))

	def test_the_database_version_is_read_with_the_engines_own_query(self):
		queries = []

		def fake_sql(query, *args, **kwargs):
			queries.append(query)
			return [("3.45.1",)] if "sqlite" in query else [("10.11.14-MariaDB",)]

		with patch.object(frappe.db, "sql", side_effect=fake_sql):
			with patch.object(frappe.db, "db_type", "sqlite"):
				self.assertEqual(database_version(), "3.45.1")
			with patch.object(frappe.db, "db_type", "mariadb"):
				self.assertEqual(database_version(), "10.11.14")
		self.assertEqual(queries, ["select sqlite_version()", "select version()"])

	def test_an_unreadable_database_version_does_not_stop_the_export(self):
		with patch.object(frappe.db, "sql", side_effect=Exception("no such function")):
			self.assertEqual(database_version(), "unknown")
		with patch.object(frappe.db, "sql", side_effect=Exception("no such function")):
			self.assertIn("unknown", generation_stamp()["database"])

	def test_a_site_without_erpnext_can_still_export(self):
		"""Global Defaults belongs to ERPNext; reading it on a Frappe-only site raises DoesNotExistError."""
		real_exists, real_single = frappe.db.exists, frappe.db.get_single_value

		def exists(doctype, name=None, *args, **kwargs):
			if doctype == "DocType" and name == "Global Defaults":
				return None
			return real_exists(doctype, name, *args, **kwargs)

		def single(doctype, *args, **kwargs):
			if doctype == "Global Defaults":
				raise frappe.DoesNotExistError("DocType Global Defaults not found")
			return real_single(doctype, *args, **kwargs)

		with (
			patch("frappe.db.exists", side_effect=exists),
			patch("frappe.db.get_single_value", side_effect=single),
		):
			self.assertEqual(default_masters()["Currency"], "USD")
			build_workbook("ToDo")  # the whole export runs without it

	# richer descriptions (batch A)

	def test_description_carries_the_metadata_that_explains_a_field(self):
		def df(**kw):
			base = dict(
				fieldname="f",
				label="F",
				fieldtype="Data",
				options=None,
				description=None,
				read_only=0,
				set_only_once=0,
				unique=0,
				fetch_from=None,
				default=None,
				mandatory_depends_on=None,
			)
			return frappe._dict({**base, **kw})

		self.assertEqual(
			describe(df(description="<b>Shown</b> on the form"), None, "X"), "F. Shown on the form."
		)
		self.assertIn("Read only.", describe(df(read_only=1), None, "X"))
		self.assertIn("Can only be set once.", describe(df(set_only_once=1), None, "X"))
		self.assertIn("Must be unique.", describe(df(unique=1), None, "X"))
		self.assertIn(
			"Fetched from customer.customer_name.",
			describe(df(fetch_from="customer.customer_name"), None, "X"),
		)
		self.assertIn("Default: 0.", describe(df(default="0"), None, "X"))
		self.assertIn("Default: today's date.", describe(df(default="Today"), None, "X"))

	@needs_erpnext
	def test_real_descriptions_now_include_read_only_and_fetch_from(self):
		by_name = {r["fieldname"]: r["description"] for r in collect_fields("Purchase Invoice")}
		self.assertIn("Read only.", by_name["base_grand_total"])
		fetched = [d for d in by_name.values() if "Fetched from" in d]
		self.assertTrue(fetched)

	@needs_erpnext
	def test_raw_expressions_are_never_shown(self):
		self.assertEqual(humanize_condition("discount"), "Required when 'discount' is set.")
		self.assertEqual(humanize_condition("eval:doc.apply_tds"), "Required when 'apply_tds' is set.")
		self.assertEqual(humanize_condition("eval:!doc.is_return"), "Required when 'is_return' is not set.")
		self.assertEqual(
			humanize_condition('eval:doc.voucher_type == "Bank Entry"'),
			"Required when 'voucher_type' is 'Bank Entry'.",
		)
		self.assertEqual(humanize_condition("eval:doc.x != 1;"), "Required when 'x' is not '1'.")
		self.assertEqual(humanize_condition("eval:doc.a && doc.b > 3"), "Required under some conditions.")
		self.assertEqual(humanize_condition("eval:frappe.user"), "Required under some conditions.")
		self.assertIsNone(humanize_condition(""))
		for description in (r["description"] for r in collect_fields("Payment Schedule", is_child=True)):
			self.assertNotIn("eval:", description)

	def test_conditionally_required_fields_are_marked_conditional(self):
		self.assertEqual(required_label({"required": True}), "Yes")
		self.assertEqual(required_label({"required": False, "conditional": True}), "Conditional")
		self.assertEqual(required_label({"required": False, "conditional": False}), "No")
		row = frappe.db.get_value(
			"DocField",
			{
				"mandatory_depends_on": ("!=", ""),
				"reqd": 0,
				"hidden": 0,
				"fieldtype": "Link",
				"parenttype": "DocType",
			},
			["parent", "fieldname"],
		)
		if not row:
			self.skipTest("no conditionally required field installed")
		field = next((r for r in collect_fields(row[0], is_child=True) if r["fieldname"] == row[1]), None)
		if field:  # skip fields the export omits (hidden or without a column)
			self.assertEqual(required_label(field), "Conditional")
			self.assertFalse(field["required"])  # a conditional field is not in the schema's required list

	# believable sample records (batch A)

	@needs_erpnext
	def test_sample_purchase_invoice_reads_like_an_unpaid_draft(self):
		record = build_payload(build_document("Purchase Invoice"))["Purchase Invoice"][0]
		item = record["items"][0]
		self.assertEqual(item["qty"] * item["rate"], item["amount"])
		self.assertEqual(record["total"], record["grand_total"])
		self.assertEqual(record["paid_amount"], 0)
		self.assertEqual(record["outstanding_amount"], record["grand_total"])
		self.assertEqual(record["total_taxes_and_charges"], 0)
		self.assertGreater(record["due_date"], record["posting_date"])
		self.assertEqual(record["conversion_rate"], 1)
		self.assertEqual(record["status"], "Draft")

	@needs_erpnext
	def test_link_examples_use_the_sites_own_settings_for_stable_masters(self):
		masters = default_masters()
		record = build_payload(build_document("Purchase Invoice"))["Purchase Invoice"][0]
		self.assertEqual(record["currency"], masters["Currency"])
		self.assertEqual(record["items"][0]["uom"], "Nos")
		self.assertEqual(record["supplier"], "Example Supplier")
		self.assertNotRegex(json.dumps(record), r"-0001")

	# a schema that can validate (batch A)

	@needs_erpnext
	def test_schema_constrains_selects_dates_and_lengths(self):
		document = build_document("Purchase Invoice")
		props = build_schema(document)["properties"]["Purchase Invoice"]["items"]["properties"]
		self.assertEqual(props["posting_date"]["format"], "date")
		self.assertIn("Draft", props["status"]["enum"])
		self.assertEqual(props["supplier"]["maxLength"], 140)
		self.assertNotIn("maxLength", props["grand_total"])  # numbers have no length
		# every example respects its own constraints (the validator checks enum, format and length)
		self.assertEqual(validate(build_payload(document), build_schema(document)), [])

	def test_select_values_allow_blank_only_when_the_options_start_blank(self):
		self.assertEqual(select_values(frappe._dict(options="\nA\nB")), ["", "A", "B"])
		self.assertEqual(select_values(frappe._dict(options="A\nB\nA")), ["A", "B"])
		self.assertEqual(select_values(frappe._dict(options="")), [])

	# version stamp (item 11)

	@needs_erpnext
	def test_schema_says_which_versions_and_database_it_came_from(self):
		schema = build_schema(build_document("Purchase Invoice"))
		stamp = schema["x-generated-from"]
		self.assertEqual(stamp["date"], frappe.utils.today())
		self.assertEqual(stamp["frappe"], frappe.__version__)
		self.assertEqual(stamp["apps"]["frappe"], frappe.__version__)
		self.assertIn("erpnext", stamp["apps"])
		self.assertRegex(stamp["database"], r"^(mariadb|postgres) \d+(\.\d+)*$")
		self.assertFalse(stamp["includes_standard_fields"])
		# the readable form carries the same facts
		self.assertIn(stamp["date"], schema["description"])
		self.assertIn(f"frappe {frappe.__version__}", schema["description"])
		self.assertIn("as reported by this database engine", schema["description"])

	@needs_erpnext
	def test_the_stamp_is_safe_to_hand_to_a_third_party(self):
		text = json.dumps(build_schema(build_document("Purchase Invoice")))
		self.assertNotIn(frappe.local.site, text)
		self.assertNotIn(frappe.session.user, json.dumps(generation_stamp()))
		self.assertEqual(
			set(generation_stamp()),
			{"date", "frappe", "apps", "database", "includes_standard_fields", "customisations_only"},
		)

	def test_an_app_without_a_version_is_reported_as_unknown(self):
		with patch("frappe.get_module", side_effect=ImportError):
			self.assertEqual(set(generation_stamp()["apps"].values()), {"unknown"})

	@needs_erpnext
	def test_the_stamp_does_not_break_validation(self):
		for doctype in ("Purchase Invoice", "ToDo"):
			self.assert_valid(doctype)

	# standard fields (item 9)

	@needs_erpnext
	def test_standard_fields_are_off_by_default(self):
		names = {r["fieldname"] for r in collect_fields("Sales Invoice")}
		self.assertFalse(names & set(STANDARD_FIELDS_PARENT))
		self.assertFalse(build_document("Sales Invoice")["include_standard"])

	@needs_erpnext
	def test_standard_fields_come_last_in_each_block(self):
		document = build_document("Sales Invoice", include_standard=True)
		self.assertEqual([f["fieldname"] for f in document["fields"]][-6:], list(STANDARD_FIELDS_PARENT))
		child = document["children"]["Sales Invoice Item"]
		self.assertEqual([f["fieldname"] for f in child][-4:], list(STANDARD_FIELDS_CHILD))
		self.assertTrue(document["include_standard"])

	@needs_erpnext
	def test_standard_fields_appear_in_the_workbook_and_stay_optional(self):
		wb, _ = self.workbook_with_standard("Sales Invoice")
		rows = {r[0]: r for r in wb["Sales Invoice"].iter_rows(min_row=2, values_only=True)}
		for name in STANDARD_FIELDS_PARENT:
			self.assertEqual(rows[name][2], "No", name)
			self.assertIn("Set by the system. Read only.", rows[name][3])
		self.assertIn("parent (child: Sales Invoice Item)", rows)
		# the required-first ordering still holds for the real fields
		names = list(rows)
		self.assertEqual(names[0], "name")

	@needs_erpnext
	def test_standard_field_types_are_the_real_database_types(self):
		types = get_column_types("Sales Invoice")
		for record in collect_fields("Sales Invoice", include_standard=True):
			if record["fieldname"] in STANDARD_FIELDS_PARENT:
				self.assertEqual(record["db_type"], types[record["fieldname"]])

	@needs_erpnext
	def test_docstatus_values_follow_whether_the_doctype_can_be_submitted(self):
		def docstatus(doctype):
			props = build_schema(build_document(doctype, include_standard=True))["properties"][doctype][
				"items"
			]["properties"]
			return props["docstatus"]

		self.assertEqual(docstatus("Sales Invoice")["enum"], [0, 1, 2])
		self.assertEqual(docstatus("ToDo")["enum"], [0])
		self.assertIn(
			"Always 0",
			{r["fieldname"]: r for r in collect_fields("ToDo", include_standard=True)}["docstatus"][
				"description"
			],
		)

	@needs_erpnext
	def test_child_rows_name_their_parent_and_table(self):
		record = build_payload(build_document("Sales Invoice", include_standard=True))["Sales Invoice"][0]
		row = record["items"][0]
		self.assertEqual(row["parenttype"], "Sales Invoice")
		self.assertEqual(row["parentfield"], "items")
		self.assertEqual(row["parent"], record["name"])
		self.assertEqual(record["docstatus"], 0)

	@needs_erpnext
	def test_payloads_with_standard_fields_still_validate(self):
		for doctype in ("Sales Invoice", "ToDo", "Contact", "Sales Invoice Item"):
			document = build_document(doctype, include_standard=True)
			self.assertEqual(validate(build_payload(document), build_schema(document)), [], doctype)

	def test_every_doctype_in_two_modules_validates_with_standard_fields(self):
		for doctype in frappe.get_all("DocType", {"module": ("in", ["Contacts", "Desk"])}, pluck="name"):
			document = build_document(doctype, include_standard=True)
			self.assertEqual(validate(build_payload(document), build_schema(document)), [], doctype)

	def test_a_single_doctype_has_no_standard_columns_to_add(self):
		names = {r["fieldname"] for r in collect_fields("System Settings", include_standard=True)}
		self.assertFalse(names & set(STANDARD_FIELDS_PARENT))

	def test_the_download_honours_the_flag_in_either_spelling(self):
		def downloaded_names(flag):
			download_workbook("ToDo", flag)
			wb = load_workbook(io.BytesIO(frappe.response["filecontent"]))
			return {r[0] for r in wb["ToDo"].iter_rows(min_row=2, values_only=True)}

		self.assertIn("docstatus", downloaded_names("1"))
		self.assertIn("docstatus", downloaded_names(1))
		self.assertNotIn("docstatus", downloaded_names("0"))
		self.assertNotIn("docstatus", downloaded_names(0))

	# preview (item 8)

	def sheet_numbers(self, doctype, include_standard):
		"""Numbers read back out of the real workbook, to compare with the preview."""
		content, _ = build_workbook(doctype, include_standard)
		rows = [r for r in load_workbook(io.BytesIO(content))[doctype].iter_rows(min_row=2, values_only=True)]
		parent = [r for r in rows if "(child:" not in r[0] and r[1] != CHILD_TABLE_TYPE]
		return {
			"rows": len(rows),
			"fields": len(parent),
			"required": sum(1 for r in parent if r[2] == "Yes"),
			"conditional": sum(1 for r in parent if r[2] == "Conditional"),
			"tables": [r[0] for r in rows if r[1] == CHILD_TABLE_TYPE],
		}

	@needs_erpnext
	def test_preview_numbers_equal_what_the_workbook_contains(self):
		for doctype in (
			"Sales Invoice",
			"Purchase Order",
			"ToDo",
			"System Settings",
			"Sales Invoice Item",
			"Contact",
		):
			for standard in (0, 1):
				preview = get_preview(doctype, standard)
				actual = self.sheet_numbers(doctype, bool(standard))
				self.assertEqual(preview["rows"], actual["rows"], (doctype, standard))
				self.assertEqual(preview["fields"], actual["fields"], (doctype, standard))
				self.assertEqual(preview["required"], actual["required"], (doctype, standard))
				self.assertEqual(preview["conditional"], actual["conditional"], (doctype, standard))
				self.assertEqual(
					[t["fieldname"] for t in preview["tables"]], actual["tables"], (doctype, standard)
				)
				self.assertEqual(preview["include_standard"], bool(standard))

	@needs_erpnext
	def test_preview_counts_each_child_tables_own_fields(self):
		tables = {t["fieldname"]: t for t in get_preview("Sales Invoice")["tables"]}
		self.assertEqual(tables["items"]["doctype"], "Sales Invoice Item")
		expected = len([r for r in collect_fields("Sales Invoice Item", is_child=True) if not r["is_table"]])
		self.assertEqual(tables["items"]["fields"], expected)

	@needs_erpnext
	def test_preview_describes_the_kind_of_doctype(self):
		self.assertEqual(get_preview("Sales Invoice")["kind"], "DocType")
		self.assertTrue(get_preview("Sales Invoice")["submittable"])
		self.assertFalse(get_preview("ToDo")["submittable"])
		self.assertEqual(get_preview("System Settings")["kind"], "Single DocType")
		self.assertEqual(get_preview("Sales Invoice Item")["kind"], "Child table")
		self.assertEqual(get_preview("Sales Invoice")["module"], "Accounts")
		custom = frappe.db.get_value("DocType", {"custom": 1, "istable": 0, "issingle": 0}, "name")
		if custom:
			self.assertTrue(get_preview(custom)["custom"])

	@needs_erpnext
	def test_preview_counts_custom_fields(self):
		if not frappe.db.exists(
			"Custom Field", {"dt": "Sales Invoice", "fieldname": "exempt_from_sales_tax"}
		):
			self.skipTest("ERPNext custom field not present")
		self.assertGreaterEqual(get_preview("Sales Invoice")["custom_fields"], 1)

	@needs_erpnext
	def test_preview_flags_a_doctype_with_nothing_to_export(self):
		preview = get_preview("Authorization Control")
		self.assertTrue(preview["empty"])
		self.assertEqual(preview["rows"], 0)
		self.assertFalse(get_preview("ToDo")["empty"])

	# friendly errors (item 10)

	@needs_erpnext
	def test_a_typed_name_is_resolved_to_the_real_one(self):
		self.assertEqual(resolve_doctype("  sales invoice "), "Sales Invoice")
		self.assertEqual(get_preview(" SALES INVOICE")["doctype"], "Sales Invoice")
		wb, filename = self.workbook(" sales invoice ")
		self.assertEqual(filename, "Sales_Invoice_Global_Schema.xlsx")
		self.assertEqual(wb.sheetnames[0], "Sales Invoice")

	def test_a_missing_doctype_is_named_in_the_error(self):
		for call in (
			lambda: build_workbook("Nonexistent Thing"),
			lambda: get_preview("Nonexistent Thing"),
			lambda: download_workbook("Nonexistent Thing"),
		):
			with self.assertRaises(frappe.ValidationError) as context:
				call()
			self.assertIn("Nonexistent Thing", str(context.exception))
			self.assertIn("does not exist", str(context.exception))

	def test_an_empty_choice_asks_for_a_doctype(self):
		for value in ("", "   ", None):
			with self.assertRaises(frappe.ValidationError) as context:
				resolve_doctype(value)
			self.assertIn("Select a DocType first", str(context.exception))

	def test_someone_without_access_is_told_who_can_export(self):
		frappe.set_user("Guest")
		in_test, frappe.flags.in_test = frappe.flags.in_test, False  # only_for skips its check in test mode
		try:
			for call in (
				lambda: get_preview("ToDo"),
				lambda: build_workbook("ToDo"),
				lambda: download_workbook("ToDo"),
			):
				with self.assertRaises(frappe.PermissionError) as context:
					call()
				self.assertIn("only allowed for", str(context.exception))
				self.assertIn("System Manager", str(context.exception))
		finally:
			frappe.flags.in_test = in_test
			frappe.clear_messages()

	# customisations only (item 12)

	def customised_names(self, doctype):
		"""Fields a Custom Field or Property Setter touches, read straight from those tables."""
		custom = set(frappe.get_all("Custom Field", {"dt": doctype}, pluck="fieldname"))
		changed = set(
			frappe.get_all(
				"Property Setter", {"doc_type": doctype, "field_name": ("is", "set")}, pluck="field_name"
			)
		)
		return custom | changed

	@needs_erpnext
	def test_customisations_only_lists_exactly_the_customised_fields(self):
		wanted = self.customised_names("Sales Invoice")
		self.assertTrue(wanted)
		everything = {r["fieldname"] for r in collect_fields("Sales Invoice") if not r["is_table"]}
		only = collect_fields("Sales Invoice", customisations_only=True)
		scalars = {r["fieldname"] for r in only if not r["is_table"]}
		self.assertEqual(scalars, wanted & everything)
		self.assertNotIn("name", scalars)
		self.assertLess(len(scalars), len(everything))

	@needs_erpnext
	def test_customisations_only_keeps_a_table_only_when_something_in_it_is_customised(self):
		document = build_document("Sales Invoice", customisations_only=True)
		for field in document["fields"]:
			if not field["is_table"]:
				continue
			child = field["child_doctype"]
			child_fields = {r["fieldname"] for r in document["children"][child]}
			itself = field["fieldname"] in self.customised_names("Sales Invoice")
			child_is_custom = bool(frappe.db.get_value("DocType", child, "custom"))
			# a table stays only for a reason...
			self.assertTrue(itself or child_fields or child_is_custom, field["fieldname"])
			# ...and its rows list nothing but customised fields (every field, if the child is custom)
			if not child_is_custom:
				self.assertLessEqual(child_fields, self.customised_names(child))
		kept = {f["child_doctype"] for f in document["fields"] if f["is_table"]}
		self.assertEqual(set(document["children"]), kept)

	def test_a_custom_doctype_is_wholly_customisation(self):
		custom = frappe.db.get_value("DocType", {"custom": 1, "istable": 0, "issingle": 0}, "name")
		if not custom:
			self.skipTest("no custom DocType installed")
		full = [r["fieldname"] for r in collect_fields(custom)]
		only = [r["fieldname"] for r in collect_fields(custom, customisations_only=True)]
		self.assertEqual(only, full)
		self.assertIn("name", only)

	def test_a_doctype_with_no_customisations_gives_an_empty_sheet(self):
		self.assertFalse(self.customised_names("ToDo"))
		preview = get_preview("ToDo", 0, 1)
		self.assertTrue(preview["empty"])
		self.assertTrue(preview["customisations_only"])
		wb, _ = self.workbook_customised("ToDo")
		self.assertEqual(wb["ToDo"].max_row, 1)  # header only

	def workbook_customised(self, doctype, include_standard=False):
		content, filename = build_workbook(doctype, include_standard, customisations_only=True)
		return load_workbook(io.BytesIO(content)), filename

	@needs_erpnext
	def test_customisations_only_says_so_in_the_schema_and_stamp(self):
		schema = build_schema(build_document("Sales Invoice", customisations_only=True))
		self.assertTrue(schema["x-generated-from"]["customisations_only"])
		self.assertIn("Only fields added or changed by customisation are listed", schema["description"])
		plain = build_schema(build_document("Sales Invoice"))
		self.assertFalse(plain["x-generated-from"]["customisations_only"])
		self.assertNotIn("Only fields added", plain["description"])

	@needs_erpnext
	def test_customisations_only_payloads_still_validate(self):
		for doctype in ("Sales Invoice", "Purchase Order", "Contact", "Address", "ToDo", "Print Settings"):
			for standard in (False, True):
				document = build_document(doctype, standard, customisations_only=True)
				self.assertEqual(
					validate(build_payload(document), build_schema(document)), [], (doctype, standard)
				)

	@needs_erpnext
	def test_customisations_only_preview_equals_the_workbook(self):
		for doctype in ("Sales Invoice", "Contact", "ToDo"):
			preview = get_preview(doctype, 0, 1)
			wb, _ = self.workbook_customised(doctype)
			rows = [r for r in wb[doctype].iter_rows(min_row=2, values_only=True)]
			self.assertEqual(preview["rows"], len(rows), doctype)

	@needs_erpnext
	def test_standard_fields_and_customisations_only_combine(self):
		document = build_document("Sales Invoice", include_standard=True, customisations_only=True)
		names = [f["fieldname"] for f in document["fields"]]
		self.assertEqual(names[-6:], list(STANDARD_FIELDS_PARENT))
		self.assertNotIn("name", names)

	@needs_erpnext
	def test_property_setters_without_a_field_are_not_counted(self):
		# DocType-level setters (no field_name) change the form, not a field
		self.assertNotIn(None, property_setter_fields("Sales Invoice"))
		self.assertNotIn("", property_setter_fields("Sales Invoice"))

	# several DocTypes at once (item 13)

	def open_pack(self, names, **flags):
		content, filename = build_pack(names, **flags)
		return zipfile.ZipFile(io.BytesIO(content)), filename

	@needs_erpnext
	def test_a_pack_holds_one_workbook_per_doctype_and_an_index(self):
		archive, filename = self.open_pack(["ToDo", "Contact", "Sales Invoice"])
		self.assertEqual(
			sorted(archive.namelist()),
			[
				"Contact_Global_Schema.xlsx",
				"Sales_Invoice_Global_Schema.xlsx",
				"ToDo_Global_Schema.xlsx",
				"_Index.xlsx",
			],
		)
		self.assertRegex(filename, r"^Fieldbook_3_DocTypes_\d{8}\.zip$")
		for name in archive.namelist():
			self.assertTrue(load_workbook(io.BytesIO(archive.read(name))).sheetnames, name)

	@needs_erpnext
	def test_the_index_matches_the_workbooks(self):
		archive, _ = self.open_pack(["ToDo", "Sales Invoice"])
		index = list(
			load_workbook(io.BytesIO(archive.read("_Index.xlsx")))["Index"].iter_rows(values_only=True)
		)
		self.assertEqual(
			index[0], ("DocType", "Module", "Kind", "Fields", "Required", "Child tables", "Rows", "File")
		)
		by_name = {row[0]: row for row in index[1:]}
		for doctype in ("ToDo", "Sales Invoice"):
			preview = get_preview(doctype)
			row = by_name[doctype]
			self.assertEqual(
				row[1:7],
				(
					preview["module"],
					preview["kind"],
					preview["fields"],
					preview["required"],
					len(preview["tables"]),
					preview["rows"],
				),
			)
			wb = load_workbook(io.BytesIO(archive.read(row[7])))
			self.assertEqual(wb[doctype].max_row - 1, row[6])

	@needs_erpnext
	def test_a_pack_follows_the_same_options_as_a_single_file(self):
		archive, _ = self.open_pack(["ToDo"], include_standard=True)
		names = {
			r[0]
			for r in load_workbook(io.BytesIO(archive.read("ToDo_Global_Schema.xlsx")))["ToDo"].iter_rows(
				min_row=2, values_only=True
			)
		}
		self.assertIn("docstatus", names)
		# compare with the same pack without the option: sites with many custom fields still shrink it
		stem = "Sales_Invoice_Global_Schema.xlsx"

		def rows(**flags):
			archive, _ = self.open_pack(["Sales Invoice"], **flags)
			return load_workbook(io.BytesIO(archive.read(stem)))["Sales Invoice"].max_row

		self.assertLess(rows(customisations_only=True), rows())

	def test_a_pack_matches_the_single_file_exports_exactly(self):
		archive, _ = self.open_pack(["Contact"])
		single, _ = build_workbook("Contact")
		packed = load_workbook(io.BytesIO(archive.read("Contact_Global_Schema.xlsx")))
		alone = load_workbook(io.BytesIO(single))
		self.assertEqual(
			[r for r in packed["Contact"].iter_rows(values_only=True)],
			[r for r in alone["Contact"].iter_rows(values_only=True)],
		)

	def test_names_are_resolved_deduplicated_and_ordered(self):
		archive, filename = self.open_pack(["todo", " ToDo ", "contact", ""])
		self.assertEqual(
			[n for n in archive.namelist() if n != "_Index.xlsx"],
			["ToDo_Global_Schema.xlsx", "Contact_Global_Schema.xlsx"],
		)
		self.assertIn("_2_DocTypes_", filename)

	def test_every_unknown_name_is_reported_together(self):
		with self.assertRaises(frappe.ValidationError) as context:
			build_pack(["ToDo", "No Such One", "Nor This"])
		message = str(context.exception)
		self.assertIn("No Such One", message)
		self.assertIn("Nor This", message)
		self.assertIn("do not exist", message)

	def test_pack_input_is_checked(self):
		for bad in ([], [""], None):
			with self.assertRaises(frappe.ValidationError) as context:
				build_pack(bad)
			self.assertIn("Select at least one DocType", str(context.exception))
		with self.assertRaises(frappe.ValidationError) as context:
			build_pack([f"DocType {i}" for i in range(MAX_PACK + 1)])
		self.assertIn(str(MAX_PACK), str(context.exception))
		with self.assertRaises(frappe.ValidationError):
			download_pack("not a list")

	def test_the_pack_download_sets_a_zip_response(self):
		download_pack(json.dumps(["ToDo", "Contact"]), 1, 0)
		self.assertEqual(frappe.response["type"], "binary")
		self.assertTrue(frappe.response["filename"].endswith(".zip"))
		self.assertTrue(zipfile.is_zipfile(io.BytesIO(frappe.response["filecontent"])))

	def test_a_pack_is_restricted_to_system_managers(self):
		frappe.set_user("Guest")
		in_test, frappe.flags.in_test = frappe.flags.in_test, False
		try:
			for call in (lambda: build_pack(["ToDo"]), lambda: download_pack(json.dumps(["ToDo"]))):
				with self.assertRaises(frappe.PermissionError) as context:
					call()
				self.assertIn("System Manager", str(context.exception))
		finally:
			frappe.flags.in_test = in_test
			frappe.clear_messages()

	# the same guarantees on any Frappe site, ERPNext or not

	def test_json_type_uses_the_fieldtype_when_the_database_says_only_text(self):
		from fieldbook.fieldbook.exporter import json_type

		# SQLite reports every column as TEXT
		self.assertEqual(json_type("TEXT", "Currency"), "number")
		self.assertEqual(json_type("TEXT", "Float"), "number")
		self.assertEqual(json_type("TEXT", "Check"), "integer")
		self.assertEqual(json_type("TEXT", "Int"), "integer")
		self.assertEqual(json_type("TEXT", "Data"), "string")
		self.assertEqual(json_type("TEXT"), "string")
		# an informative database type still wins
		self.assertEqual(json_type("int(11)", "Data"), "integer")
		self.assertEqual(json_type("decimal(21,9)", "Data"), "number")

	def test_core_doctypes_validate_in_every_option_combination(self):
		for doctype in ("Contact", "ToDo", "User", "System Settings", "DocType", "Role", "Notification"):
			for standard in (False, True):
				for customised in (False, True):
					document = build_document(doctype, standard, customised)
					self.assertEqual(
						validate(build_payload(document), build_schema(document)),
						[],
						(doctype, standard, customised),
					)

	def test_numbers_are_numbers_in_the_payload_and_the_schema(self):
		# a DocType with Currency, Float, Int and Check fields; wrong on engines that report TEXT for all
		document = build_document("Dashboard Chart")
		props = build_schema(document)["properties"]["Dashboard Chart"]["items"]["properties"]
		record = build_payload(document)["Dashboard Chart"][0]
		numeric = [
			f for f in document["fields"] if f["df"].fieldtype in ("Int", "Check", "Float", "Currency")
		]
		self.assertTrue(numeric)
		for field in numeric:
			expected = "integer" if field["df"].fieldtype in ("Int", "Check") else "number"
			self.assertEqual(props[field["fieldname"]]["type"], expected, field["fieldname"])
			self.assertNotIsInstance(record[field["fieldname"]], str, field["fieldname"])

	def test_core_workbook_layout_and_ordering(self):
		wb, filename = self.workbook("Contact")
		self.assertEqual(wb.sheetnames, ["Contact", "Payload", "Schema"])
		self.assertEqual(filename, "Contact_Global_Schema.xlsx")
		sheet = wb["Contact"]
		self.assertEqual(
			[c.value for c in sheet[1]], ["Field Name", "Data Type", "Required", "Description", "Example"]
		)
		rows = [[c.value for c in r] for r in sheet.iter_rows(min_row=2)]
		self.assertEqual(rows[0][0], "name")
		parent = [r for r in rows if "(child:" not in r[0] and r[1] != CHILD_TABLE_TYPE]
		flags = [r[2] == "Yes" for r in parent]
		self.assertEqual(flags, sorted(flags, reverse=True))

	def test_core_child_tables_are_flattened(self):
		names = [r[0] for r in self.workbook("Contact")[0]["Contact"].iter_rows(min_row=2, values_only=True)]
		at = names.index("email_ids")
		self.assertEqual(names[at + 1], "email_id (child: Contact Email)")

	def test_core_database_types_are_what_the_database_reports(self):
		types = get_column_types("Contact")
		for row in self.workbook("Contact")[0]["Contact"].iter_rows(min_row=2, values_only=True):
			if "(child:" in row[0] or row[1] == CHILD_TABLE_TYPE or row[0] not in types:
				continue
			self.assertEqual(row[1], types[row[0]], row[0])

	def test_core_stamp_and_standard_fields(self):
		schema = build_schema(build_document("Contact", include_standard=True))
		stamp = schema["x-generated-from"]
		self.assertEqual(stamp["frappe"], frappe.__version__)
		self.assertTrue(stamp["database"].startswith(frappe.db.db_type + " "))
		self.assertNotIn(frappe.local.site, json.dumps(schema))
		document = build_document("Contact", include_standard=True)
		if get_column_types("Contact"):  # engines that can describe the table
			self.assertEqual([f["fieldname"] for f in document["fields"]][-6:], list(STANDARD_FIELDS_PARENT))

	def test_core_customisations_are_a_subset_of_the_full_export(self):
		for doctype in ("Contact", "ToDo", "User"):
			full = {f["fieldname"] for f in collect_fields(doctype) if not f["is_table"]}
			only = {
				f["fieldname"] for f in collect_fields(doctype, customisations_only=True) if not f["is_table"]
			}
			self.assertLessEqual(only, full | {"name"}, doctype)
			self.assertLessEqual(only - {"name"}, self.customised_names(doctype), doctype)

	def test_core_preview_equals_the_workbook(self):
		for doctype in ("Contact", "ToDo", "System Settings"):
			for standard in (0, 1):
				preview = get_preview(doctype, standard)
				actual = self.sheet_numbers(doctype, bool(standard))
				self.assertEqual(preview["rows"], actual["rows"], (doctype, standard))
				self.assertEqual(preview["fields"], actual["fields"], (doctype, standard))

	def test_core_pack_has_an_index_that_matches(self):
		archive, _ = self.open_pack(["ToDo", "Contact", "Role"])
		rows = list(
			load_workbook(io.BytesIO(archive.read("_Index.xlsx")))["Index"].iter_rows(values_only=True)
		)
		self.assertEqual([r[0] for r in rows[1:]], ["ToDo", "Contact", "Role"])
		for row in rows[1:]:
			wb = load_workbook(io.BytesIO(archive.read(row[7])))
			self.assertEqual(wb[row[0]].max_row - 1, row[6])

	# every kind of DocType

	def test_single_doctype_has_fields_but_no_name_row(self):
		document = self.assert_valid("System Settings")
		by_name = {f["fieldname"]: f for f in document["fields"]}
		self.assertNotIn("name", by_name)
		# no table to read, so types come from Frappe's own fieldtype map
		language = by_name["language"]
		self.assertEqual(language["db_type"], type_from_fieldtype(language["df"].fieldtype))

	def test_virtual_doctype_exports(self):
		virtual = frappe.get_all("DocType", {"is_virtual": 1, "istable": 0}, pluck="name")
		if not virtual:
			self.skipTest("no virtual DocType installed")
		for doctype in virtual:
			self.assert_valid(doctype)

	def test_auto_increment_doctype_has_an_integer_name(self):
		doctype = frappe.db.get_value("DocType", {"autoname": "autoincrement", "istable": 0}, "name")
		if not doctype:
			self.skipTest("no autoincrement DocType installed")
		name = next(f for f in self.assert_valid(doctype)["fields"] if f["fieldname"] == "name")
		self.assertEqual(name["example"], 1)

	def test_custom_doctypes_export(self):
		custom = frappe.get_all("DocType", {"custom": 1}, pluck="name")
		if not custom:
			self.skipTest("no custom DocType installed")
		for doctype in custom:
			self.assert_valid(doctype)

	@needs_erpnext
	def test_a_child_table_exports_on_its_own(self):
		self.assert_valid("Sales Invoice Item")

	@needs_erpnext
	def test_doctype_with_no_fields_still_exports_a_valid_workbook(self):
		wb, _ = self.workbook("Authorization Control")
		self.assertEqual(wb["Authorization Control"].max_row, 1)  # header only

	def test_every_doctype_in_two_whole_modules_exports_and_validates(self):
		names = frappe.get_all("DocType", {"module": ("in", ["Contacts", "Desk"])}, pluck="name")
		self.assertGreater(len(names), 30)
		for doctype in names:
			self.assert_valid(doctype)

	def test_illegal_characters_and_oversized_text_are_made_safe(self):
		self.assertEqual(safe_cell("a\x00b\x0bc"), "abc")
		self.assertEqual(len(safe_cell("x" * 40000)), 32000)
		self.assertEqual(safe_cell(5), 5)

	# access and the page

	def test_export_is_restricted_to_system_managers(self):
		# frappe.only_for skips its check while Frappe is in test mode, so turn that off for this
		# test to exercise what a real request does
		frappe.set_user("Guest")
		in_test, frappe.flags.in_test = frappe.flags.in_test, False
		try:
			with self.assertRaises(frappe.PermissionError):
				build_workbook("ToDo")
			with self.assertRaises(frappe.PermissionError):
				download_workbook("ToDo")
		finally:
			frappe.flags.in_test = in_test

	def test_export_of_a_missing_doctype_fails(self):
		with self.assertRaises(frappe.ValidationError):
			build_workbook("No Such DocType")

	def test_download_sets_a_binary_response(self):
		download_workbook("ToDo")
		self.assertEqual(frappe.response["type"], "binary")
		self.assertEqual(frappe.response["filename"], "ToDo_Global_Schema.xlsx")
		self.assertTrue(load_workbook(io.BytesIO(frappe.response["filecontent"])).sheetnames)

	def test_the_desk_page_is_registered_for_system_managers(self):
		page = frappe.get_doc("Page", "fieldbook")
		self.assertEqual(page.title, "Fieldbook")
		self.assertEqual([r.role for r in page.roles], ["System Manager"])
		# the page's script downloads through this method, so it must stay whitelisted
		method = frappe.get_attr("fieldbook.fieldbook.exporter.download_workbook")
		self.assertIn(method, frappe.whitelisted)
